"""Microsoft PowerPoint (.pptx) extractor."""

from pathlib import Path
from typing import Optional, Dict, Any, List
import pptx
from src.ingestion.extractors.base import BaseExtractor, ExtractedDocument, ExtractedBlock
from src.logging_config import rag_logger


class PptxExtractor(BaseExtractor):
    """Extracts slides, slide titles, body bullet points, and tables from PowerPoint files."""

    def extract(self, file_path: Path, metadata: Optional[Dict[str, Any]] = None) -> ExtractedDocument:
        metadata = metadata or {}
        prs = pptx.Presentation(file_path)
        blocks: List[ExtractedBlock] = []

        rag_logger.info(f"Extracting PPTX: {file_path.name} ({len(prs.slides)} slides)")

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = f"Slide {slide_idx}"
            slide_texts: List[str] = []

            # Check title shape
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text and text != slide_title:
                            slide_texts.append(text)

                elif shape.has_table:
                    table = shape.table
                    rows_data = []
                    for row in table.rows:
                        rows_data.append([c.text.strip().replace("\n", " ") for c in row.cells])
                    if rows_data:
                        headers = rows_data[0]
                        sep = ["---"] * len(headers)
                        table_lines = [
                            f"[Slide {slide_idx} Table: {slide_title}]",
                            "| " + " | ".join(headers) + " |",
                            "| " + " | ".join(sep) + " |"
                        ]
                        for r in rows_data[1:]:
                            table_lines.append("| " + " | ".join(r) + " |")

                        blocks.append(ExtractedBlock(
                            text="\n".join(table_lines),
                            page_number=slide_idx,
                            block_type="table",
                            section=slide_title,
                            is_table=True
                        ))

            # Combine slide body text
            if slide_texts:
                full_slide_text = f"## {slide_title}\n" + "\n".join(f"- {t}" for t in slide_texts)
                blocks.append(ExtractedBlock(
                    text=full_slide_text,
                    page_number=slide_idx,
                    block_type="paragraph",
                    section=slide_title
                ))

            # Slide notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                note_text = slide.notes_slide.notes_text_frame.text.strip()
                if note_text:
                    blocks.append(ExtractedBlock(
                        text=f"[Presenter Notes for Slide {slide_idx}]:\n{note_text}",
                        page_number=slide_idx,
                        block_type="paragraph",
                        section=slide_title
                    ))

        return ExtractedDocument(
            filename=file_path.name,
            file_type="pptx",
            total_pages=len(prs.slides),
            blocks=blocks,
            metadata=metadata
        )
