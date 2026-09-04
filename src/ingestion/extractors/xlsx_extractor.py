"""Microsoft Excel (.xlsx) extractor."""

from pathlib import Path
from typing import Optional, Dict, Any, List
import openpyxl
from src.ingestion.extractors.base import BaseExtractor, ExtractedDocument, ExtractedBlock
from src.logging_config import rag_logger


class XlsxExtractor(BaseExtractor):
    """Extracts tabular records and equipment logs from Excel spreadsheets."""

    def extract(self, file_path: Path, metadata: Optional[Dict[str, Any]] = None) -> ExtractedDocument:
        metadata = metadata or {}
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        blocks: List[ExtractedBlock] = []
        page_idx = 1

        rag_logger.info(f"Extracting XLSX: {file_path.name} (Sheets: {wb.sheetnames})")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Filter out empty rows
            non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() != "" for cell in r)]
            if not non_empty_rows:
                continue

            # First row is typically header
            headers = [str(c or "").strip() for c in non_empty_rows[0]]

            # Split large sheets into manageable tabular blocks (e.g. 20 rows per chunk)
            chunk_size = 20
            data_rows = non_empty_rows[1:]

            if not data_rows:
                # Only header or single row
                line = " | ".join(headers)
                blocks.append(ExtractedBlock(
                    text=f"[Sheet: {sheet_name}]\n{line}",
                    page_number=page_idx,
                    block_type="table",
                    section=f"Sheet: {sheet_name}",
                    is_table=True
                ))
            else:
                for i in range(0, len(data_rows), chunk_size):
                    batch = data_rows[i:i + chunk_size]
                    sep = ["---"] * len(headers)
                    lines = [
                        f"[Sheet: {sheet_name} | Rows {i+1}-{i+len(batch)}]",
                        "| " + " | ".join(headers) + " |",
                        "| " + " | ".join(sep) + " |"
                    ]
                    for r in batch:
                        row_vals = [str(c or "").strip().replace("\n", " ") for c in r]
                        # Pad or trim to match headers length
                        if len(row_vals) < len(headers):
                            row_vals += [""] * (len(headers) - len(row_vals))
                        lines.append("| " + " | ".join(row_vals[:len(headers)]) + " |")

                    table_text = "\n".join(lines)
                    blocks.append(ExtractedBlock(
                        text=table_text,
                        page_number=page_idx,
                        block_type="table",
                        section=f"Sheet: {sheet_name}",
                        is_table=True
                    ))

            page_idx += 1

        wb.close()
        return ExtractedDocument(
            filename=file_path.name,
            file_type="xlsx",
            total_pages=max(1, page_idx - 1),
            blocks=blocks,
            metadata=metadata
        )
